from io import BytesIO

from docx import Document as WordDocument
from fastapi.testclient import TestClient

from app.database import Membership, Tenant, User
from app.security.passwords import hash_password


def _case(client: TestClient, auth: dict[str, str]) -> str:
    response = client.post(
        "/api/v1/cases",
        headers=auth,
        json={"title": "合同字段抽取", "description": ""},
    )
    return response.json()["id"]


def _contract_docx() -> bytes:
    document = WordDocument()
    document.add_paragraph("采购服务合同")
    document.add_paragraph("合同编号：HT-2026-001")
    document.add_paragraph("甲方：示例采购有限公司")
    document.add_paragraph("乙方：示例供应商有限公司")
    document.add_paragraph("合同总额：人民币 12,345.60")
    document.add_paragraph("生效日期：2026年9月27日")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_docx_classification_fields_and_evidence(client: TestClient, auth):
    case_id = _case(client, auth)
    upload = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers=auth,
        files={
            "file": (
                "contract.docx",
                _contract_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201, upload.text
    document_id = upload.json()["id"]
    submitted = client.post(
        f"/api/v1/cases/{case_id}/submit",
        headers={**auth, "Idempotency-Key": "docx-extraction"},
    )
    assert submitted.status_code == 202, submitted.text

    detail = client.get(f"/api/v1/cases/{case_id}", headers=auth).json()
    parsed = detail["documents"][0]
    assert parsed["parse_status"] == "COMPLETED"
    assert parsed["classification"] == "contract"
    assert parsed["page_count"] == 1

    response = client.get(f"/api/v1/documents/{document_id}/extraction", headers=auth)
    assert response.status_code == 200, response.text
    result = response.json()
    fields = {field["field_name"]: field for field in result["fields"]}
    assert fields["contract_number"]["normalized_value"] == "HT-2026-001"
    assert fields["total_amount"]["normalized_value"] == "12345.60"
    assert fields["effective_date"]["normalized_value"] == "2026-09-27"
    assert fields["total_amount"]["evidence"]["page_number"] == 1
    assert len(fields["total_amount"]["evidence"]["text_hash"]) == 64


def test_extraction_is_tenant_scoped(client: TestClient, auth, pdf_bytes, database):
    case_id = _case(client, auth)
    upload = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers=auth,
        files={"file": ("contract.pdf", pdf_bytes, "application/pdf")},
    )
    document_id = upload.json()["id"]
    client.post(
        f"/api/v1/cases/{case_id}/submit",
        headers={**auth, "Idempotency-Key": "tenant-extraction"},
    )
    assert client.get(f"/api/v1/documents/{document_id}/extraction", headers=auth).status_code == 200
    with database.session_factory() as session:
        tenant = Tenant(name="Extraction Isolation Tenant")
        user = User(
            email="extract-other@example.com",
            display_name="Other",
            password_hash=hash_password("Other123!"),
        )
        session.add_all([tenant, user])
        session.flush()
        session.add(Membership(tenant_id=tenant.id, user_id=user.id, roles="reviewer"))
        session.commit()
    login = client.post(
        "/api/v1/auth/login",
        json={"email": "extract-other@example.com", "password": "Other123!"},
    )
    other_auth = {"Authorization": f"Bearer {login.json()['access_token']}"}
    assert (
        client.get(f"/api/v1/documents/{document_id}/extraction", headers=other_auth).status_code
        == 404
    )


def test_corrupt_pdf_records_failed_state(client: TestClient, auth):
    case_id = _case(client, auth)
    upload = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers=auth,
        files={"file": ("broken.pdf", b"%PDF-this-is-not-valid", "application/pdf")},
    )
    assert upload.status_code == 201
    submitted = client.post(
        f"/api/v1/cases/{case_id}/submit",
        headers={**auth, "Idempotency-Key": "broken-pdf"},
    )
    job = client.get(f"/api/v1/jobs/{submitted.json()['id']}", headers=auth).json()
    assert job["status"] == "FAILED"
    assert job["error_code"] == "ParseError"
    case = client.get(f"/api/v1/cases/{case_id}", headers=auth).json()
    assert case["status"] == "FAILED"
    assert case["documents"][0]["parse_status"] == "FAILED"
