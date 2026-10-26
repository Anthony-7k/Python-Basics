from io import BytesIO

from docx import Document as WordDocument
from fastapi.testclient import TestClient

from app.database import Membership, ReviewCase, Tenant, User
from app.security.passwords import hash_password


def _docx() -> bytes:
    document = WordDocument()
    document.add_paragraph("采购合同")
    document.add_paragraph("乙方：示例供应商有限公司")
    document.add_paragraph("合同总额：10000.00")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_dashboard_summarizes_latest_tenant_results(client: TestClient, auth, database):
    draft = client.post(
        "/api/v1/cases",
        headers=auth,
        json={"title": "草稿案件", "description": ""},
    )
    assert draft.status_code == 201
    processed = client.post(
        "/api/v1/cases",
        headers=auth,
        json={"title": "已处理案件", "description": ""},
    ).json()
    upload = client.post(
        f"/api/v1/cases/{processed['id']}/documents",
        headers=auth,
        files={
            "file": (
                "contract.docx",
                _docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    client.post(
        f"/api/v1/cases/{processed['id']}/submit",
        headers={**auth, "Idempotency-Key": "dashboard-job"},
    )
    extraction = client.get(
        f"/api/v1/documents/{upload.json()['id']}/extraction", headers=auth
    ).json()
    field = next(item for item in extraction["fields"] if item["field_name"] == "party_b")
    corrected = client.patch(
        f"/api/v1/fields/{field['id']}",
        headers=auth,
        json={
            "expected_version": field["version"],
            "normalized_value": "示例供应商集团有限公司",
            "reason": "仪表盘人工修正样例",
        },
    )
    assert corrected.status_code == 200

    with database.session_factory() as session:
        tenant = Tenant(name="Dashboard Isolation Tenant")
        user = User(
            email="dashboard-other@example.com",
            display_name="Other",
            password_hash=hash_password("Other123!"),
        )
        session.add_all([tenant, user])
        session.flush()
        session.add(Membership(tenant_id=tenant.id, user_id=user.id, roles="submitter"))
        session.add(
            ReviewCase(
                tenant_id=tenant.id,
                owner_id=user.id,
                title="其他租户案件",
                description="不应进入统计",
            )
        )
        session.commit()

    response = client.get("/api/v1/dashboard/summary", headers=auth)
    assert response.status_code == 200
    summary = response.json()
    assert summary["total_cases"] == 2
    assert summary["status_counts"] == {"DRAFT": 1, "REVIEW_READY": 1}
    assert summary["total_documents"] == 1
    assert summary["total_jobs"] == 1
    assert summary["job_success_rate"] == 100.0
    assert summary["average_processing_seconds"] >= 0
    assert summary["manual_review_rate"] == 50.0
    assert summary["rule_hit_rate"] == 100.0
    assert summary["risk_by_severity"] == {"HIGH": 1}
    assert summary["risk_by_status"] == {"OPEN": 1}
    assert summary["approval_counts"] == {}
    assert summary["ready_reports"] == 0


def test_dashboard_date_filters_and_validation(client, auth):
    empty = client.get(
        "/api/v1/dashboard/summary",
        headers=auth,
        params={"start": "2099-01-01", "end": "2099-12-31"},
    )
    assert empty.status_code == 200
    assert empty.json()["total_cases"] == 0
    invalid = client.get(
        "/api/v1/dashboard/summary",
        headers=auth,
        params={"start": "2026-12-31", "end": "2026-01-01"},
    )
    assert invalid.status_code == 400
