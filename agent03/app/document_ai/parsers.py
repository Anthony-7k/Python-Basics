from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

import fitz
from docx import Document as WordDocument
from PIL import Image, ImageOps

from app.domain import DocumentKind

MAX_PAGES = 20


class ParseError(ValueError):
    pass


@dataclass(frozen=True)
class TextBlock:
    text: str
    bbox: list[float]
    confidence: float = 0.98


@dataclass(frozen=True)
class ParsedPage:
    page_number: int
    text: str
    width: float
    height: float
    blocks: list[TextBlock]
    method: str = "native"


class OCRProvider(Protocol):
    def recognize(self, payload: bytes) -> list[TextBlock]: ...


def parse_document(path: Path, kind: DocumentKind, original_name: str,
                   ocr: OCRProvider | None = None) -> list[ParsedPage]:
    if kind == DocumentKind.PDF:
        return _parse_pdf(path, ocr)
    if kind == DocumentKind.DOCX:
        return _parse_docx(path)
    if kind in {DocumentKind.PNG, DocumentKind.JPEG}:
        return _parse_image(path, original_name, ocr)
    raise ParseError(f"unsupported document kind: {kind}")


def _parse_pdf(path: Path, ocr: OCRProvider | None = None) -> list[ParsedPage]:
    try:
        pdf = fitz.open(stream=path.read_bytes(), filetype="pdf")
    except Exception as exc:
        raise ParseError("PDF cannot be opened") from exc
    with pdf:
        if pdf.page_count > MAX_PAGES:
            raise ParseError(f"document exceeds the {MAX_PAGES}-page limit")
        pages = []
        for index, page in enumerate(pdf):
            method = "native"
            blocks = [
                TextBlock(text=item[4].strip(), bbox=[item[0], item[1], item[2], item[3]])
                for item in page.get_text("blocks")
                if item[4].strip()
            ]
            if ocr is not None and len("".join(block.text for block in blocks).strip()) < 20:
                scale = min(2.5, (20_000_000 / max(1, page.rect.width * page.rect.height)) ** 0.5)
                pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                if pixmap.width * pixmap.height > 20_000_000:
                    raise ParseError("PDF raster exceeds OCR pixel limit")
                recognized = ocr.recognize(pixmap.tobytes("png"))
                if not recognized:
                    raise ParseError(f"OCR found no text on page {index + 1}")
                blocks = [TextBlock(block.text,
                                    [block.bbox[0] * page.rect.width / pixmap.width,
                                     block.bbox[1] * page.rect.height / pixmap.height,
                                     block.bbox[2] * page.rect.width / pixmap.width,
                                     block.bbox[3] * page.rect.height / pixmap.height],
                                    block.confidence) for block in recognized]
                method = "rapidocr"
            pages.append(
                ParsedPage(
                    page_number=index + 1,
                    text="\n".join(block.text for block in blocks),
                    width=float(page.rect.width),
                    height=float(page.rect.height),
                    blocks=blocks,
                    method=method,
                )
            )
        return pages


def _parse_docx(path: Path) -> list[ParsedPage]:
    try:
        document = WordDocument(path)
    except Exception as exc:
        raise ParseError("DOCX cannot be opened") from exc
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    blocks = [TextBlock(text=line, bbox=[0.0, float(i), 1.0, float(i + 1)]) for i, line in enumerate(lines)]
    return [
        ParsedPage(
            page_number=1,
            text="\n".join(lines),
            width=1.0,
            height=max(1.0, float(len(lines))),
            blocks=blocks,
        )
    ]


def _parse_image(path: Path, original_name: str, ocr: OCRProvider | None = None) -> list[ParsedPage]:
    try:
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image)
            width, height = image.size
            if width * height > 20_000_000:
                raise ParseError("image exceeds OCR pixel limit")
            if ocr is not None:
                from io import BytesIO

                buffer = BytesIO()
                image.convert("RGB").save(buffer, format="PNG")
                blocks = ocr.recognize(buffer.getvalue())
                if not blocks:
                    raise ParseError("OCR found no text in image")
                return [ParsedPage(1, "\n".join(block.text for block in blocks),
                                   float(width), float(height), blocks, "rapidocr")]
    except Exception as exc:
        raise ParseError("image cannot be opened") from exc
    text = f"[Fake OCR pending] {original_name}"
    return [
        ParsedPage(
            page_number=1,
            text=text,
            width=float(width),
            height=float(height),
            blocks=[TextBlock(text=text, bbox=[0.0, 0.0, float(width), float(height)])],
            method="fake",
        )
    ]

